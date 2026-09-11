"""
Company AI Excel Batch Tool (Tkinter)

功能：
1. 从 Excel 指定列逐行读取文本；
2. 将“固定 Prompt + 当前行文本”发送给公司 AI；
3. 将 AI 返回结果写入同一行的指定输出列；
4. 每完成一行立即保存，可中断后继续；
5. Token 仅在界面中临时输入，不写入代码或配置文件；
6. Prompt 输入框留空时自动使用内置 DEFAULT_PROMPT；
7. 固定使用 V4（ds_v4）、standard agent、Thinking 关闭、SSL 验证关闭；
8. 固定 4 并发：最多同时处理 4 行，每行使用独立 session；
9. Prompt 支持占位符：写在 Prompt 里的 {employee_prompt} 会被替换成当前行的
   员工 prompt；也支持 <<<EMPLOYEE_PROMPT_START ... EMPLOYEE_PROMPT_END>>>
   标记替换。两种方式都没有时，才退化为把文本追加到 Prompt 末尾；
10. 固定运行 2 次：每行调用 2 次 AI，2 次结果分别写入“结果列1 / 结果列2”。

依赖：
    pip install requests openpyxl

说明：
- 默认接口地址根据提供的测试代码整理，请以浏览器 Inspect/Network 中的真实地址为准。
- 请确认自动调用内部 API 符合公司的授权、信息安全和使用频率要求。
"""

from __future__ import annotations

import json
import queue
from concurrent.futures import FIRST_COMPLETED, ThreadPoolExecutor, wait
import re
import threading
import time
import traceback
from dataclasses import dataclass
from pathlib import Path
from typing import Any, Callable, Iterator

import openpyxl
import requests
import tkinter as tk
import urllib3
from openpyxl.styles import Alignment
from openpyxl.utils import column_index_from_string, get_column_letter
from tkinter import filedialog, messagebox, scrolledtext, ttk


# -----------------------------------------------------------------------------
# 默认 Prompt：如果界面的“固定 Prompt”输入框留空，程序自动使用下面内容。
# 你之后可以直接修改这里的文字，替换成公司最终确认的 Prompt。
# -----------------------------------------------------------------------------
DEFAULT_PROMPT = """请检查下面的文本中是否包含敏感信息。
请识别并列出所有可能的敏感信息，并简要说明其敏感类型。
如果未发现敏感信息，请明确回答“未发现敏感信息”。
只根据给定文本进行判断，不要补充文本中不存在的信息。

待检测员工prompt：
<<<EMPLOYEE_PROMPT_START
{employee_prompt}
EMPLOYEE_PROMPT_END>>>"""

# 员工 prompt 占位符：出现在 Prompt 中时，会被当前行的员工 prompt 替换。
EMPLOYEE_PROMPT_PLACEHOLDER = "{employee_prompt}"
# 备用替换方式：若 Prompt 中没有占位符，但包含下面这对标记，
# 则把两个标记之间的内容替换成员工 prompt。
EMPLOYEE_PROMPT_START = "<<<EMPLOYEE_PROMPT_START"
EMPLOYEE_PROMPT_END = "EMPLOYEE_PROMPT_END>>>"

# 根据参考代码整理的默认配置；若 Network 中地址不同，请在界面修改。
DEFAULT_SESSION_URL = (
    "https://audit-ai-api.apps.ocp-dta-sh.cn.kworld.kpmg.com/api/session"
)
DEFAULT_CHAT_URL = (
    "https://audit-ai-api.apps.ocp-dta-sh.cn.kworld.kpmg.com/api/chat"
)
DEFAULT_AGENT_ID = "standard"
DEFAULT_MODEL_NAME = "ds_v4"

# 固定并发数：最多同时向公司 AI 处理 4 行。
MAX_CONCURRENT_WORKERS = 4

# 固定运行次数：每一行都对 AI 发起 2 次调用，结果分别写入 2 个结果列。
# 修改此值时必须保证 UI 中的结果列数量与之匹配（见 output_columns）。
RUN_COUNT = 2

# 每个并发 worker 完成一行后固定冷却 1 秒，再处理下一行。
# 该参数不在 UI 中暴露。
DEFAULT_DELAY_SECONDS = 1.0

# 内部站点若使用公司自签名证书，通常需要关闭 SSL 验证。
urllib3.disable_warnings(urllib3.exceptions.InsecureRequestWarning)


@dataclass(frozen=True)
class ApiConfig:
    token: str
    session_url: str
    chat_url: str
    agent_id: str
    model_name: str
    stream: bool = True
    enable_thinking: bool = False
    verify_ssl: bool = False
    connect_timeout: int = 30
    read_timeout: int = 300


class CompanyAIClient:
    """调用公司 AI session/chat 接口，并解析 SSE 流式返回。"""

    def __init__(self, config: ApiConfig):
        self.config = config
        self.http = requests.Session()

    @staticmethod
    def normalize_token(token: str) -> str:
        token = token.strip()
        if not token:
            raise ValueError("Authorization Token 为空。")
        if not token.lower().startswith("bearer "):
            token = f"Bearer {token}"
        return token

    def _headers(self, *, sse: bool = False) -> dict[str, str]:
        return {
            "Content-Type": "application/json",
            "Accept": "text/event-stream" if sse else "application/json",
            "Accept-Encoding": "gzip, deflate, br",
            "Authorization": self.normalize_token(self.config.token),
        }

    def create_session(self) -> str:
        response = self.http.post(
            self.config.session_url,
            json={"agent_id": self.config.agent_id},
            headers=self._headers(),
            verify=self.config.verify_ssl,
            timeout=(self.config.connect_timeout, self.config.read_timeout),
        )
        self._raise_for_status_with_body(response)

        try:
            result = response.json()
        except ValueError as exc:
            raise RuntimeError(
                f"创建会话接口没有返回 JSON：{response.text[:500]}"
            ) from exc

        # 参考代码格式：{"success": true, "data": {"id": "..."}}
        if result.get("success") is False:
            raise RuntimeError(
                str(result.get("ret_msg") or result.get("message") or result)
            )

        data = result.get("data") or {}
        session_id = data.get("id") or result.get("session_id") or result.get("id")
        if not session_id:
            raise RuntimeError(f"创建会话成功，但响应中找不到 session id：{result}")
        return str(session_id)

    def ask(self, prompt: str) -> str:
        """每次调用新建一个 session，然后发送一条完整 prompt。"""
        session_id = self.create_session()
        return self.send_in_session(session_id, prompt)

    def send_in_session(self, session_id: str, prompt: str) -> str:
        payload = {
            "agent_id": self.config.agent_id,
            "session_id": session_id,
            "message": prompt,
            "config": {
                "model_name": self.config.model_name,
                "stream": self.config.stream,
                "enable_thinking": self.config.enable_thinking,
            },
        }

        response = self.http.post(
            self.config.chat_url,
            json=payload,
            headers=self._headers(sse=self.config.stream),
            stream=self.config.stream,
            verify=self.config.verify_ssl,
            timeout=(self.config.connect_timeout, self.config.read_timeout),
        )
        self._raise_for_status_with_body(response)

        try:
            if not self.config.stream:
                return self._parse_non_stream_response(response)
            return self._parse_stream_response(response)
        finally:
            response.close()

    @staticmethod
    def _raise_for_status_with_body(response: requests.Response) -> None:
        if response.ok:
            return
        body = ""
        try:
            body = response.text[:1000]
        except Exception:
            pass
        raise requests.HTTPError(
            f"HTTP {response.status_code} {response.reason}; response={body}",
            response=response,
        )

    def _parse_non_stream_response(self, response: requests.Response) -> str:
        try:
            obj = response.json()
        except ValueError:
            return response.text.strip()

        fragments = list(self._extract_ai_text_fragments(obj))
        if fragments:
            return self._merge_fragments(fragments).strip()
        return self._extract_text(obj).strip()

    def _parse_stream_response(self, response: requests.Response) -> str:
        answer = ""
        saw_json_event = False

        for event_type, data_text in self._iter_sse_events(response):
            if data_text in {"", "[DONE]"}:
                if data_text == "[DONE]":
                    break
                continue

            # 参考代码只处理 updates；这里也兼容未提供 event 类型的 SSE。
            if event_type and event_type not in {"updates", "message", "data"}:
                continue
            if data_text.startswith("Not supported event"):
                continue

            try:
                obj = json.loads(data_text)
            except json.JSONDecodeError:
                # 某些接口可能直接以 data: 文本片段返回。
                answer = self._merge_piece(answer, data_text)
                continue

            saw_json_event = True
            for fragment in self._extract_ai_text_fragments(obj):
                answer = self._merge_piece(answer, fragment)

        # 如果响应不是标准 SSE，尝试读取普通响应正文。
        if not answer and not saw_json_event:
            raw = response.text.strip()
            if raw:
                try:
                    obj = json.loads(raw)
                    fragments = list(self._extract_ai_text_fragments(obj))
                    answer = self._merge_fragments(fragments)
                except json.JSONDecodeError:
                    answer = raw

        return answer.strip()

    @staticmethod
    def _iter_sse_events(
        response: requests.Response,
    ) -> Iterator[tuple[str | None, str]]:
        """逐个产生 (event_type, data)，兼容多行 data。"""
        event_type: str | None = None
        data_lines: list[str] = []

        for raw_line in response.iter_lines(decode_unicode=True):
            if raw_line is None:
                continue
            line = raw_line if isinstance(raw_line, str) else raw_line.decode(
                "utf-8", errors="ignore"
            )

            if line == "":
                if data_lines:
                    yield event_type, "\n".join(data_lines).strip()
                event_type = None
                data_lines = []
                continue

            if line.startswith(":"):
                continue  # SSE heartbeat/comment
            if line.startswith("event:"):
                event_type = line[6:].strip() or None
            elif line.startswith("data:"):
                data_lines.append(line[5:].lstrip())

        if data_lines:
            yield event_type, "\n".join(data_lines).strip()

    @classmethod
    def _extract_ai_text_fragments(cls, obj: Any) -> Iterator[str]:
        """兼容参考代码中的 data.messages / messages / 单消息结构。"""
        if not isinstance(obj, dict):
            text = cls._extract_text(obj)
            if text:
                yield text
            return

        data = obj.get("data")
        if isinstance(data, dict) and isinstance(data.get("messages"), list):
            messages = data["messages"]
        elif isinstance(obj.get("messages"), list):
            messages = obj["messages"]
        else:
            messages = [obj]

        yielded = False
        for message in messages:
            if not isinstance(message, dict):
                continue

            message_type = message.get("type")
            if message_type is None:
                if message.get("tool_call_id") or message.get("name") == "query_kb":
                    continue
            elif str(message_type).lower() not in {
                "ai",
                "assistant",
                "answer",
            }:
                continue

            content = message.get("content")
            if content is None:
                content = message.get("text")
            fragment = cls._extract_text(content)
            if fragment:
                yielded = True
                yield fragment

        # 部分接口可能直接返回 answer/output/content。
        if not yielded:
            for key in ("answer", "output", "content", "text", "value"):
                if key in obj:
                    fragment = cls._extract_text(obj[key])
                    if fragment:
                        yield fragment
                        return

    @classmethod
    def _extract_text(cls, value: Any) -> str:
        if value is None:
            return ""
        if isinstance(value, str):
            return value
        if isinstance(value, (int, float, bool)):
            return str(value)
        if isinstance(value, list):
            return "".join(cls._extract_text(item) for item in value)
        if isinstance(value, dict):
            for key in ("text", "content", "value", "answer", "output"):
                if key in value:
                    return cls._extract_text(value[key])
            return ""
        return str(value)

    @staticmethod
    def _merge_piece(existing: str, fragment: str) -> str:
        """兼容“增量片段”和“每次返回累计全文”两种 SSE 格式。"""
        fragment = fragment or ""
        if not fragment:
            return existing
        if not existing:
            return fragment
        if fragment == existing or existing.endswith(fragment):
            return existing
        if fragment.startswith(existing):
            return fragment

        # 找 existing 后缀与 fragment 前缀的最大重叠，避免重复拼接。
        max_overlap = min(len(existing), len(fragment))
        for size in range(max_overlap, 0, -1):
            if existing[-size:] == fragment[:size]:
                return existing + fragment[size:]
        return existing + fragment

    @classmethod
    def _merge_fragments(cls, fragments: list[str]) -> str:
        answer = ""
        for fragment in fragments:
            answer = cls._merge_piece(answer, fragment)
        return answer


@dataclass(frozen=True)
class BatchOptions:
    input_file: Path
    output_file: Path
    sheet_name: str
    input_column: int
    # 结果列：每一列对应一次 AI 调用（默认 2 列 = 运行 2 次）。
    output_columns: tuple[int, ...]
    start_row: int
    prompt_template: str
    delay_seconds: float
    skip_filled: bool


def build_prompt(prompt_template: str, row_text: str) -> str:
    """
    把当前行的员工 prompt 组装进固定 Prompt。

    优先级：
    1. Prompt 中含 {employee_prompt} → 用员工 prompt 原地替换该占位符；
    2. Prompt 中含 <<<EMPLOYEE_PROMPT_START ... EMPLOYEE_PROMPT_END>>> →
       用员工 prompt 替换两个标记之间的内容；
    3. 两者都没有 → 退化为把员工 prompt 追加到 Prompt 末尾（旧行为）。

    注意：这里不使用 str.format()，避免 Prompt 或员工文本中出现的
    大括号 {} 触发格式化异常。
    """
    template = (prompt_template or "").strip()
    text = row_text.strip()
    if not template:
        return text

    # 情况 1：占位符替换（最常用）。
    if EMPLOYEE_PROMPT_PLACEHOLDER in template:
        return template.replace(EMPLOYEE_PROMPT_PLACEHOLDER, text)

    # 情况 2：只有成对标记、没有占位符时，替换标记之间的内容。
    start_at = template.find(EMPLOYEE_PROMPT_START)
    if start_at != -1:
        content_start = start_at + len(EMPLOYEE_PROMPT_START)
        end_at = template.find(EMPLOYEE_PROMPT_END, content_start)
        if end_at != -1:
            return (
                template[:content_start]
                + "\n"
                + text
                + "\n"
                + template[end_at:]
            )

    # 情况 3：兜底，保持旧的追加行为。
    return f"{template}\n\n【待检查文本】\n{text}"


def _run_one_ai_task(
    api_config: ApiConfig,
    prompt: str,
    delay_seconds: float,
    run_count: int,
) -> list[tuple[bool, str]]:
    """
    单个并发任务：独立 HTTP client，按 run_count 次调用 AI。

    每一次调用都会新建独立 session（client.ask 内部完成），
    因此同一行的两次结果互不共享上下文，便于对比两次输出是否一致。

    返回 [(是否成功, 文本), ...]，长度等于 run_count，顺序与结果列一一对应。
    """
    client = CompanyAIClient(api_config)
    outcomes: list[tuple[bool, str]] = []
    try:
        for _ in range(max(1, run_count)):
            try:
                answer = client.ask(prompt)
                if not answer:
                    raise RuntimeError("接口调用成功，但没有解析到 AI 回答。")
                outcomes.append((True, answer))
            except Exception as exc:
                outcomes.append((False, f"{type(exc).__name__}: {exc}"))
    finally:
        try:
            client.http.close()
        finally:
            if delay_seconds > 0:
                time.sleep(delay_seconds)
    return outcomes


def process_excel(
    api_config: ApiConfig,
    options: BatchOptions,
    stop_event: threading.Event,
    on_log: Callable[[str], None],
    on_progress: Callable[[int, int], None],
) -> dict[str, int]:
    """
    固定 4 并发处理 Excel；每行固定运行 RUN_COUNT(=2) 次，结果分别写入 2 个结果列。

    设计原则：
    - API 请求由最多 4 个 worker 并发执行；
    - 每行独立创建 CompanyAIClient 和 session，避免上下文互相污染；
    - 同一行内部的多次调用串行执行，且每次都会新建独立 session；
    - openpyxl 的写入和 workbook.save() 只在当前批处理线程中执行，
      不让 4 个 API worker 同时写同一个 Excel 文件。
    """
    input_suffix = options.input_file.suffix.lower()
    if input_suffix not in {".xlsx", ".xlsm"}:
        raise ValueError("仅支持 .xlsx 或 .xlsm 文件；旧版 .xls 请先另存为 .xlsx。")

    keep_vba = input_suffix == ".xlsm"
    workbook = openpyxl.load_workbook(
        options.input_file,
        keep_vba=keep_vba,
        data_only=False,
    )
    if options.sheet_name not in workbook.sheetnames:
        workbook.close()
        raise ValueError(f"工作表不存在：{options.sheet_name}")
    worksheet = workbook[options.sheet_name]

    rows: list[int] = []
    for row in range(options.start_row, worksheet.max_row + 1):
        value = worksheet.cell(row, options.input_column).value
        if value is not None and str(value).strip():
            rows.append(row)

    if not rows:
        workbook.close()
        raise ValueError(
            f"工作表 {options.sheet_name} 的 "
            f"{get_column_letter(options.input_column)} 列没有可处理文字。"
        )

    options.output_file.parent.mkdir(parents=True, exist_ok=True)
    try:
        workbook.save(options.output_file)
    except PermissionError as exc:
        workbook.close()
        raise PermissionError(
            f"无法写入 {options.output_file.name}。请先关闭已打开的 Excel 文件。"
        ) from exc

    total = len(rows)
    success_count = 0
    error_count = 0
    skipped_count = 0
    completed_count = 0

    # work_items: (行号, 员工文本, 需要写入的结果列下标元组)
    # 下标元组用于跳过已有结果的列，从而支持中断后续跑。
    work_items: list[tuple[int, str, tuple[int, ...]]] = []
    for row in rows:
        missing_targets = tuple(
            index
            for index, column in enumerate(options.output_columns)
            if worksheet.cell(row, column).value in (None, "")
        )
        if options.skip_filled and not missing_targets:
            skipped_count += 1
            completed_count += 1
            on_log(
                f"第 {row} 行：{len(options.output_columns)} 个结果列均已有内容，跳过。"
            )
            on_progress(completed_count, total)
            continue

        if options.skip_filled:
            # 只补跑缺失的结果列（例如上次被中断）。
            targets = missing_targets
        else:
            # 不跳过时整行重跑，覆盖所有结果列。
            targets = tuple(range(len(options.output_columns)))

        row_text = str(worksheet.cell(row, options.input_column).value).strip()
        work_items.append((row, row_text, targets))

    if not work_items:
        workbook.save(options.output_file)
        workbook.close()
        return {
            "total": total,
            "success": success_count,
            "errors": error_count,
            "skipped": skipped_count,
        }

    on_log(
        f"固定 {MAX_CONCURRENT_WORKERS} 并发启动：最多同时处理 "
        f"{MAX_CONCURRENT_WORKERS} 行；每行运行 {RUN_COUNT} 次，结果分别写入 "
        f"{' / '.join(get_column_letter(col) for col in options.output_columns)} 列。"
    )

    # 只维持最多 4 个 in-flight future，不一次性把所有行都排入线程池。
    # 这样点击“停止”后，只需等待当前最多 4 个请求完成，不会继续启动后续行。
    item_iter = iter(work_items)
    pending: dict[Any, tuple[int, str, tuple[int, ...]]] = {}

    def submit_next(executor: ThreadPoolExecutor) -> bool:
        if stop_event.is_set():
            return False
        try:
            row, row_text, targets = next(item_iter)
        except StopIteration:
            return False

        full_prompt = build_prompt(options.prompt_template, row_text)
        on_log(
            f"第 {row} 行：已提交 AI（{len(targets)} 次调用；{row_text[:50]}"
            f"{'…' if len(row_text) > 50 else ''}）"
        )
        future = executor.submit(
            _run_one_ai_task,
            api_config,
            full_prompt,
            options.delay_seconds,
            len(targets),
        )
        pending[future] = (row, row_text, targets)
        return True

    try:
        with ThreadPoolExecutor(
            max_workers=MAX_CONCURRENT_WORKERS,
            thread_name_prefix="CompanyAI",
        ) as executor:
            for _ in range(MAX_CONCURRENT_WORKERS):
                if not submit_next(executor):
                    break

            while pending:
                done_futures, _ = wait(
                    tuple(pending.keys()),
                    return_when=FIRST_COMPLETED,
                )

                for future in done_futures:
                    row, _row_text, targets = pending.pop(future)

                    try:
                        outcomes = future.result()
                    except Exception as exc:
                        # 任务整体异常时，所有目标列都写失败信息。
                        outcomes = [
                            (False, f"{type(exc).__name__}: {exc}")
                            for _ in targets
                        ]

                    ok_values: list[str] = []
                    for target_index, (ok, value) in zip(targets, outcomes):
                        column = options.output_columns[target_index]
                        output_cell = worksheet.cell(row, column)
                        output_cell.value = value if ok else f"ERROR: {value}"
                        output_cell.alignment = Alignment(
                            wrap_text=True,
                            vertical="top",
                        )
                        if ok:
                            success_count += 1
                            ok_values.append(value)
                        else:
                            error_count += 1
                            on_log(
                                f"第 {row} 行 / {get_column_letter(column)} 列："
                                f"失败：{value}"
                            )

                    if len(ok_values) == len(targets):
                        on_log(
                            f"第 {row} 行：{len(targets)} 次调用全部完成（{ok_values[0][:60]}"
                            f"{'…' if len(ok_values[0]) > 60 else ''}）"
                        )
                    else:
                        on_log(
                            f"第 {row} 行：{len(ok_values)}/{len(targets)} 次调用成功，"
                            "失败内容已写入对应单元格。"
                        )

                    # Excel 永远只在此线程写入/保存，不在 API worker 中写。
                    try:
                        workbook.save(options.output_file)
                    except PermissionError as exc:
                        raise PermissionError(
                            f"保存失败：{options.output_file.name} 正被 Excel 占用，"
                            "请关闭该文件。"
                        ) from exc

                    completed_count += 1
                    on_progress(completed_count, total)

                    if not stop_event.is_set():
                        submit_next(executor)

                if stop_event.is_set() and pending:
                    on_log(
                        f"已停止提交新任务；等待当前 {len(pending)} 个并发请求完成。"
                    )

        if stop_event.is_set():
            remaining = total - completed_count
            if remaining > 0:
                on_log(
                    f"停止完成：尚有 {remaining} 行未处理，可下次继续运行。"
                )

        workbook.save(options.output_file)
        return {
            "total": total,
            "success": success_count,
            "errors": error_count,
            "skipped": skipped_count,
        }
    finally:
        workbook.close()


class ExcelAIToolApp(tk.Tk):
    def __init__(self) -> None:
        super().__init__()
        self.title("Company AI Excel Batch Tool - 4 Concurrent / 2 Runs")
        self.geometry("920x750")
        self.minsize(820, 670)

        self.ui_queue: queue.Queue[tuple[str, Any]] = queue.Queue()
        self.stop_event = threading.Event()
        self.worker: threading.Thread | None = None

        self._create_variables()
        self._build_ui()
        self.after(100, self._poll_queue)

    def _create_variables(self) -> None:
        self.input_file_var = tk.StringVar()
        self.output_file_var = tk.StringVar()
        self.sheet_var = tk.StringVar()
        self.input_col_var = tk.StringVar(value="A")
        self.output_col1_var = tk.StringVar(value="B")
        self.output_col2_var = tk.StringVar(value="C")
        self.start_row_var = tk.StringVar(value="2")

        self.token_var = tk.StringVar()
        self.session_url_var = tk.StringVar(value=DEFAULT_SESSION_URL)
        self.chat_url_var = tk.StringVar(value=DEFAULT_CHAT_URL)
        self.skip_filled_var = tk.BooleanVar(value=True)
        self.show_token_var = tk.BooleanVar(value=False)
        self.status_var = tk.StringVar(value="请选择 Excel，粘贴 Token，然后开始运行。")

    def _build_ui(self) -> None:
        main = ttk.Frame(self, padding=12)
        main.pack(fill="both", expand=True)
        main.columnconfigure(1, weight=1)

        row = 0
        ttk.Label(main, text="输入 Excel：").grid(row=row, column=0, sticky="w", pady=4)
        ttk.Entry(main, textvariable=self.input_file_var).grid(
            row=row, column=1, sticky="ew", padx=6
        )
        ttk.Button(main, text="选择…", command=self._choose_input_file).grid(
            row=row, column=2, sticky="ew"
        )

        row += 1
        ttk.Label(main, text="结果 Excel：").grid(row=row, column=0, sticky="w", pady=4)
        ttk.Entry(main, textvariable=self.output_file_var).grid(
            row=row, column=1, sticky="ew", padx=6
        )
        ttk.Button(main, text="另存为…", command=self._choose_output_file).grid(
            row=row, column=2, sticky="ew"
        )

        row += 1
        excel_options = ttk.Frame(main)
        excel_options.grid(row=row, column=0, columnspan=3, sticky="ew", pady=(4, 8))
        for col in (1, 3, 5, 7, 9):
            excel_options.columnconfigure(col, weight=1)
        ttk.Label(excel_options, text="工作表").grid(row=0, column=0, padx=(0, 4))
        self.sheet_combo = ttk.Combobox(
            excel_options, textvariable=self.sheet_var, state="readonly", width=18
        )
        self.sheet_combo.grid(row=0, column=1, sticky="ew", padx=(0, 12))
        ttk.Label(excel_options, text="文字列").grid(row=0, column=2, padx=(0, 4))
        ttk.Entry(excel_options, textvariable=self.input_col_var, width=7).grid(
            row=0, column=3, sticky="ew", padx=(0, 12)
        )
        ttk.Label(excel_options, text="结果列1").grid(row=0, column=4, padx=(0, 4))
        ttk.Entry(excel_options, textvariable=self.output_col1_var, width=7).grid(
            row=0, column=5, sticky="ew", padx=(0, 12)
        )
        ttk.Label(excel_options, text="结果列2").grid(row=0, column=6, padx=(0, 4))
        ttk.Entry(excel_options, textvariable=self.output_col2_var, width=7).grid(
            row=0, column=7, sticky="ew", padx=(0, 12)
        )
        ttk.Label(excel_options, text="起始行").grid(row=0, column=8, padx=(0, 4))
        ttk.Entry(excel_options, textvariable=self.start_row_var, width=7).grid(
            row=0, column=9, sticky="ew"
        )

        row += 1
        ttk.Separator(main).grid(row=row, column=0, columnspan=3, sticky="ew", pady=8)

        row += 1
        ttk.Label(main, text="Authorization：").grid(
            row=row, column=0, sticky="w", pady=4
        )
        self.token_entry = ttk.Entry(main, textvariable=self.token_var, show="●")
        self.token_entry.grid(row=row, column=1, sticky="ew", padx=6)
        ttk.Checkbutton(
            main,
            text="显示",
            variable=self.show_token_var,
            command=self._toggle_token_visibility,
        ).grid(row=row, column=2, sticky="w")

        row += 1
        ttk.Label(main, text="Session URL：").grid(row=row, column=0, sticky="w", pady=4)
        ttk.Entry(main, textvariable=self.session_url_var).grid(
            row=row, column=1, columnspan=2, sticky="ew", padx=6
        )

        row += 1
        ttk.Label(main, text="Chat URL：").grid(row=row, column=0, sticky="w", pady=4)
        ttk.Entry(main, textvariable=self.chat_url_var).grid(
            row=row, column=1, columnspan=2, sticky="ew", padx=6
        )

        row += 1
        check_frame = ttk.Frame(main)
        check_frame.grid(row=row, column=0, columnspan=3, sticky="w", pady=4)
        ttk.Checkbutton(
            check_frame, text="跳过已有结果的单元格", variable=self.skip_filled_var
        ).pack(side="left")

        row += 1
        ttk.Label(main, text="固定 Prompt：").grid(
            row=row, column=0, sticky="nw", pady=(8, 4)
        )
        self.prompt_text = scrolledtext.ScrolledText(main, height=9, wrap="word")
        self.prompt_text.grid(
            row=row, column=1, columnspan=2, sticky="nsew", padx=6, pady=(8, 4)
        )
        # 默认保持空白：留空时运行阶段自动使用 DEFAULT_PROMPT。
        # 这样用户可以直接运行默认规则，也可以在此临时输入自定义 Prompt。
        main.rowconfigure(row, weight=1)

        row += 1
        ttk.Label(
            main,
            text=(
                "提示：Prompt 留空时自动使用内置 DEFAULT_PROMPT。"
                "Prompt 中写 {employee_prompt} 会替换成当前行的员工 prompt；"
                "也可用 <<<EMPLOYEE_PROMPT_START ... EMPLOYEE_PROMPT_END>>> 包裹。"
            ),
        ).grid(row=row, column=1, columnspan=2, sticky="w", padx=6, pady=(0, 4))

        row += 1
        button_frame = ttk.Frame(main)
        button_frame.grid(row=row, column=0, columnspan=3, sticky="ew", pady=8)
        self.start_button = ttk.Button(
            button_frame, text="开始批量检查", command=self._start
        )
        self.start_button.pack(side="left")
        self.stop_button = ttk.Button(
            button_frame, text="停止（完成当前并发请求后）", command=self._stop, state="disabled"
        )
        self.stop_button.pack(side="left", padx=8)
        ttk.Button(
            button_frame, text="打开结果所在文件夹", command=self._open_output_folder
        ).pack(side="left")

        row += 1
        self.progress = ttk.Progressbar(main, mode="determinate")
        self.progress.grid(row=row, column=0, columnspan=3, sticky="ew", pady=(0, 4))

        row += 1
        ttk.Label(main, textvariable=self.status_var).grid(
            row=row, column=0, columnspan=3, sticky="w", pady=(0, 4)
        )

        row += 1
        self.log_text = scrolledtext.ScrolledText(
            main, height=10, wrap="word", state="disabled"
        )
        self.log_text.grid(row=row, column=0, columnspan=3, sticky="nsew")
        main.rowconfigure(row, weight=1)

    def _toggle_token_visibility(self) -> None:
        self.token_entry.configure(show="" if self.show_token_var.get() else "●")

    def _choose_input_file(self) -> None:
        file_path = filedialog.askopenfilename(
            title="选择 Excel 文件",
            filetypes=[
                ("Excel workbook", "*.xlsx *.xlsm"),
                ("All files", "*.*"),
            ],
        )
        if not file_path:
            return
        self.input_file_var.set(file_path)
        path = Path(file_path)
        suffix = path.suffix
        self.output_file_var.set(str(path.with_name(f"{path.stem}_AI_result{suffix}")))
        self._load_sheet_names(path)

    def _choose_output_file(self) -> None:
        initial = Path(self.output_file_var.get()) if self.output_file_var.get() else None
        file_path = filedialog.asksaveasfilename(
            title="保存结果 Excel",
            defaultextension=".xlsx",
            initialdir=str(initial.parent) if initial else None,
            initialfile=initial.name if initial else "AI_result.xlsx",
            filetypes=[
                ("Excel workbook", "*.xlsx"),
                ("Macro-enabled workbook", "*.xlsm"),
            ],
        )
        if file_path:
            self.output_file_var.set(file_path)

    def _load_sheet_names(self, path: Path) -> None:
        try:
            workbook = openpyxl.load_workbook(path, read_only=True)
            names = workbook.sheetnames
            workbook.close()
            self.sheet_combo["values"] = names
            if names:
                self.sheet_var.set(names[0])
        except Exception as exc:
            messagebox.showerror("读取 Excel 失败", str(exc))

    def _validate_and_build(self) -> tuple[ApiConfig, BatchOptions]:
        input_file = Path(self.input_file_var.get().strip())
        output_file = Path(self.output_file_var.get().strip())
        if not input_file.exists():
            raise ValueError("请选择存在的输入 Excel 文件。")
        if not output_file.name:
            raise ValueError("请设置结果 Excel 文件。")
        if input_file.resolve() == output_file.resolve():
            raise ValueError("结果文件请使用另一个文件名，以保护原始 Excel。")

        token = self.token_var.get().strip()
        if not token:
            raise ValueError(
                "请粘贴 Inspect → Network → app_info → Authorization 中的 Bearer Token。"
            )

        session_url = self.session_url_var.get().strip()
        chat_url = self.chat_url_var.get().strip()
        if not session_url.startswith(("http://", "https://")):
            raise ValueError("Session URL 格式不正确。")
        if not chat_url.startswith(("http://", "https://")):
            raise ValueError("Chat URL 格式不正确。")

        input_col_text = self.input_col_var.get().strip().upper()
        output_col1_text = self.output_col1_var.get().strip().upper()
        output_col2_text = self.output_col2_var.get().strip().upper()
        if not re.fullmatch(r"[A-Z]{1,3}", input_col_text):
            raise ValueError("文字列请输入 Excel 列字母，例如 A。")
        if not re.fullmatch(r"[A-Z]{1,3}", output_col1_text):
            raise ValueError("结果列1 请输入 Excel 列字母，例如 B。")
        if not re.fullmatch(r"[A-Z]{1,3}", output_col2_text):
            raise ValueError("结果列2 请输入 Excel 列字母，例如 C。")
        input_column = column_index_from_string(input_col_text)
        output_columns = (
            column_index_from_string(output_col1_text),
            column_index_from_string(output_col2_text),
        )
        if len(output_columns) != RUN_COUNT:
            raise ValueError(
                f"结果列数量必须等于运行次数 RUN_COUNT={RUN_COUNT}。"
            )
        if output_columns[0] == output_columns[1]:
            raise ValueError("结果列1 和结果列2 不能相同。")
        if input_column in output_columns:
            raise ValueError("文字列不能与任一结果列相同。")

        try:
            start_row = int(self.start_row_var.get())
            if start_row < 1:
                raise ValueError
        except ValueError as exc:
            raise ValueError("起始行必须是大于等于 1 的整数。") from exc

        delay_seconds = DEFAULT_DELAY_SECONDS

        sheet_name = self.sheet_var.get().strip()
        if not sheet_name:
            raise ValueError("请选择工作表。")

        # Agent / Model / Thinking / SSL 均固定在代码中，不在 UI 暴露。
        # 与当前测试环境保持一致：standard + V4(ds_v4) + thinking=False + SSL verify=False。
        api_config = ApiConfig(
            token=token,
            session_url=session_url,
            chat_url=chat_url,
            agent_id=DEFAULT_AGENT_ID,
            model_name=DEFAULT_MODEL_NAME,
            stream=True,
            enable_thinking=False,
            verify_ssl=False,
        )

        typed_prompt = self.prompt_text.get("1.0", "end-1c").strip()
        effective_prompt = typed_prompt if typed_prompt else DEFAULT_PROMPT.strip()

        batch_options = BatchOptions(
            input_file=input_file,
            output_file=output_file,
            sheet_name=sheet_name,
            input_column=input_column,
            output_columns=output_columns,
            start_row=start_row,
            prompt_template=effective_prompt,
            delay_seconds=delay_seconds,
            skip_filled=self.skip_filled_var.get(),
        )
        return api_config, batch_options

    def _start(self) -> None:
        if self.worker and self.worker.is_alive():
            return
        try:
            api_config, batch_options = self._validate_and_build()
        except Exception as exc:
            messagebox.showerror("配置错误", str(exc))
            return

        self.stop_event.clear()
        self.progress["value"] = 0
        self.status_var.set("正在运行…")
        self.start_button.configure(state="disabled")
        self.stop_button.configure(state="normal")
        self._append_log("=" * 60)
        self._append_log(f"输入文件：{batch_options.input_file}")
        self._append_log(f"结果文件：{batch_options.output_file}")
        self._append_log(
            f"工作表：{batch_options.sheet_name}；"
            f"文字列：{get_column_letter(batch_options.input_column)}；"
            f"结果列："
            f"{' / '.join(get_column_letter(col) for col in batch_options.output_columns)}"
        )
        self._append_log(f"Model：V4 ({api_config.model_name})")
        self._append_log(f"并发数：{MAX_CONCURRENT_WORKERS}（固定）")
        self._append_log(
            f"每行运行次数：{RUN_COUNT}（固定，结果分别写入各结果列）"
        )

        typed_prompt = self.prompt_text.get("1.0", "end-1c").strip()
        current_prompt = typed_prompt or DEFAULT_PROMPT.strip()
        if typed_prompt:
            self._append_log("Prompt：使用界面中输入的自定义 Prompt。")
        else:
            self._append_log("Prompt：输入框为空，使用内置 DEFAULT_PROMPT。")
        if EMPLOYEE_PROMPT_PLACEHOLDER in current_prompt:
            self._append_log(
                f"Prompt：检测到 {EMPLOYEE_PROMPT_PLACEHOLDER} 占位符，"
                "员工 prompt 会插入该位置。"
            )
        elif (
            EMPLOYEE_PROMPT_START in current_prompt
            and EMPLOYEE_PROMPT_END in current_prompt
        ):
            self._append_log(
                "Prompt：检测到员工 prompt 标记，员工 prompt 会插入标记之间。"
            )
        else:
            self._append_log(
                "Prompt：未检测到占位符/标记，员工 prompt 将追加到 Prompt 末尾。"
            )

        self.worker = threading.Thread(
            target=self._worker_main,
            args=(api_config, batch_options),
            daemon=True,
        )
        self.worker.start()

    def _worker_main(self, api_config: ApiConfig, options: BatchOptions) -> None:
        try:
            summary = process_excel(
                api_config=api_config,
                options=options,
                stop_event=self.stop_event,
                on_log=lambda text: self.ui_queue.put(("log", text)),
                on_progress=lambda done, total: self.ui_queue.put(
                    ("progress", (done, total))
                ),
            )
            self.ui_queue.put(("done", (summary, str(options.output_file))))
        except Exception as exc:
            self.ui_queue.put(
                (
                    "fatal",
                    (
                        f"{type(exc).__name__}: {exc}",
                        traceback.format_exc(),
                    ),
                )
            )

    def _stop(self) -> None:
        self.stop_event.set()
        self.status_var.set("已请求停止；不再提交新行，等待当前最多 4 个请求结束。")
        self._append_log("用户请求停止。")
        self.stop_button.configure(state="disabled")

    def _poll_queue(self) -> None:
        try:
            while True:
                event, payload = self.ui_queue.get_nowait()
                if event == "log":
                    self._append_log(str(payload))
                elif event == "progress":
                    done, total = payload
                    self.progress["maximum"] = max(total, 1)
                    self.progress["value"] = done
                    self.status_var.set(f"处理中：{done}/{total}")
                elif event == "done":
                    summary, output_file = payload
                    self._finish_running_state()
                    stopped = self.stop_event.is_set()
                    status = "已停止并保存" if stopped else "处理完成"
                    self.status_var.set(
                        f"{status}：共 {summary['total']} 行；"
                        f"成功 {summary['success']} 次，"
                        f"失败 {summary['errors']} 次，"
                        f"跳过 {summary['skipped']} 行。"
                    )
                    self._append_log(self.status_var.get())
                    messagebox.showinfo(
                        status,
                        f"共处理行数：{summary['total']}\n"
                        f"成功：{summary['success']} 次\n"
                        f"失败：{summary['errors']} 次\n"
                        f"跳过：{summary['skipped']} 行\n\n"
                        f"结果文件：\n{output_file}",
                    )
                elif event == "fatal":
                    error, details = payload
                    self._finish_running_state()
                    self.status_var.set(f"运行失败：{error}")
                    self._append_log(details)
                    messagebox.showerror("运行失败", error)
        except queue.Empty:
            pass
        finally:
            self.after(100, self._poll_queue)

    def _finish_running_state(self) -> None:
        self.start_button.configure(state="normal")
        self.stop_button.configure(state="disabled")

    def _append_log(self, text: str) -> None:
        self.log_text.configure(state="normal")
        self.log_text.insert("end", f"{time.strftime('%H:%M:%S')}  {text}\n")
        self.log_text.see("end")
        self.log_text.configure(state="disabled")

    def _open_output_folder(self) -> None:
        output = self.output_file_var.get().strip()
        folder = Path(output).parent if output else Path.cwd()
        try:
            import os
            import subprocess
            import sys

            if sys.platform.startswith("win"):
                os.startfile(folder)  # type: ignore[attr-defined]
            elif sys.platform == "darwin":
                subprocess.run(["open", str(folder)], check=False)
            else:
                subprocess.run(["xdg-open", str(folder)], check=False)
        except Exception as exc:
            messagebox.showerror("无法打开文件夹", str(exc))


if __name__ == "__main__":
    app = ExcelAIToolApp()
    app.mainloop()
